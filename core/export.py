"""Markdown → Word(.docx) / PDF 変換（議事録をそのまま提出できる形式に）"""
import re
from pathlib import Path

RED_SPAN_RE = re.compile(r'<span style="color:red">(.*?)</span>', re.DOTALL)
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
INLINE_SPLIT_RE = re.compile(r'(<span style="color:red">.*?</span>|\*\*.+?\*\*)', re.DOTALL)


def _table_cells(line):
    return [cell.strip().replace(r"\|", "|")
            for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def _blocks(text):
    """Markdown表は段落と区別する。セル数の不足は空欄で補う。"""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        if i + 1 < len(lines) and "|" in lines[i]:
            header = _table_cells(lines[i])
            separator = _table_cells(lines[i + 1])
            if len(header) == len(separator) and all(re.fullmatch(r":?-{3,}:?", c) for c in separator):
                rows = [header]
                i += 2
                while i < len(lines) and "|" in lines[i] and lines[i].strip():
                    cells = _table_cells(lines[i])
                    # 余分なセルを黙って落とさない。
                    width = max(len(rows[0]), len(cells))
                    rows = [r + [""] * (width - len(r)) for r in rows]
                    rows.append(cells + [""] * (width - len(cells)))
                    i += 1
                yield "table", rows
                continue
        yield _parse_line(lines[i])
        i += 1


def _parse_line(line: str):
    """行を (kind, content) に分類する。
    kind: h1/h2/h3, check, bullet, text, blank"""
    stripped = line.rstrip()
    if not stripped.strip():
        return "blank", ""
    m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if m:
        level = min(len(m.group(1)), 3)
        return f"h{level}", m.group(2)
    m = re.match(r"^\s*-\s*\[[ xX]?\]\s*(.*)$", stripped)
    if m:
        return "check", m.group(1)
    m = re.match(r"^\s*[-*]\s+(.*)$", stripped)
    if m:
        return "bullet", m.group(1)
    return "text", stripped


def _inline_runs(text: str):
    """インライン書式を (text, is_red, is_bold) のランに分解する"""
    runs = []
    for part in INLINE_SPLIT_RE.split(text):
        if not part:
            continue
        red_m = RED_SPAN_RE.fullmatch(part)
        if red_m:
            inner = red_m.group(1)
            bold_inner = BOLD_RE.fullmatch(inner)
            if bold_inner:
                runs.append((bold_inner.group(1), True, True))
            else:
                runs.append((inner, True, False))
            continue
        bold_m = BOLD_RE.fullmatch(part)
        if bold_m:
            runs.append((bold_m.group(1), False, True))
            continue
        runs.append((part, False, False))
    return runs


def md_to_docx(md_path: str, out_path: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    text = Path(md_path).read_text(encoding="utf-8")

    for kind, content in _blocks(text):
        if kind == "table":
            table = doc.add_table(rows=0, cols=len(content[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(content):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    _add_runs_docx(cell.paragraphs[0], re.sub(r"<br\s*/?>", "\n", value))
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
            continue
        if kind == "blank":
            continue
        if kind.startswith("h"):
            level = int(kind[1])
            para = doc.add_heading("", level=level)
            _add_runs_docx(para, content)
        elif kind == "check":
            para = doc.add_paragraph(style="List Bullet")
            para.add_run("☐ ")
            _add_runs_docx(para, content)
        elif kind == "bullet":
            para = doc.add_paragraph(style="List Bullet")
            _add_runs_docx(para, content)
        else:
            para = doc.add_paragraph()
            _add_runs_docx(para, content)

    doc.save(out_path)
    return out_path


def _add_runs_docx(para, content: str):
    from docx.shared import RGBColor
    for text, is_red, is_bold in _inline_runs(content):
        run = para.add_run(text)
        if is_red:
            run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        if is_bold:
            run.bold = True


def md_to_pdf(md_path: str, out_path: str) -> str:
    """reportlab の日本語CIDフォント（HeiseiKakuGo-W5）でPDF化。
    フォントファイル不要で Windows / Mac 両対応。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, LongTable, TableStyle

    font_name = "HeiseiKakuGo-W5"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:
        font_name = "Helvetica"  # 最終フォールバック（日本語は出ない）

    styles = {
        "h1": ParagraphStyle("h1", fontName=font_name, fontSize=16, leading=22, spaceAfter=6),
        "h2": ParagraphStyle("h2", fontName=font_name, fontSize=13, leading=19, spaceAfter=5),
        "h3": ParagraphStyle("h3", fontName=font_name, fontSize=12, leading=18, spaceAfter=4),
        "text": ParagraphStyle("text", fontName=font_name, fontSize=10, leading=16),
        "bullet": ParagraphStyle("bullet", fontName=font_name, fontSize=10, leading=16,
                                 leftIndent=12),
    }

    text = Path(md_path).read_text(encoding="utf-8")
    flow = []
    for kind, content in _blocks(text):
        if kind == "table":
            rows = [[Paragraph(_inline_markup_pdf(value), styles["text"]) for value in row] for row in content]
            table = LongTable(rows, colWidths=[(A4[0] - 36 * mm) / len(rows[0])] * len(rows[0]),
                              repeatRows=1, splitInRow=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            flow.extend([table, Spacer(1, 3 * mm)])
            continue
        if kind == "blank":
            flow.append(Spacer(1, 3 * mm))
            continue
        markup = _inline_markup_pdf(content)
        if kind.startswith("h"):
            flow.append(Paragraph(markup, styles[kind]))
        elif kind == "check":
            flow.append(Paragraph("☐ " + markup, styles["bullet"]))
        elif kind == "bullet":
            flow.append(Paragraph("・" + markup, styles["bullet"]))
        else:
            flow.append(Paragraph(markup, styles["text"]))

    doc = SimpleDocTemplate(out_path, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    doc.build(flow)
    return out_path


def _inline_markup_pdf(content: str) -> str:
    """reportlab Paragraph 用の <font color> / <b> マークアップに変換"""
    from xml.sax.saxutils import escape
    parts = []
    for text, is_red, is_bold in _inline_runs(content):
        t = escape(text)
        if is_bold:
            t = f"<b>{t}</b>"
        if is_red:
            t = f'<font color="#cc0000">{t}</font>'
        parts.append(t)
    return "".join(parts)
