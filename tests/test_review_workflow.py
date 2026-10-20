"""編集の喪失・校正漏れ・表の崩れを、個人データを使わずに再現する。"""
import copy
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import Mock, patch
from types import SimpleNamespace

import app
from core import config, document_status, export, job_history, minutes, output, pipeline, search
from core import utterances as U


def utterance(text="開始は10時です。", start=0):
    return dict(start=start, end=start + 3, text=text, marked=text, speaker="話者1", words=[])


class ReviewWorkflowTest(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.folder = self.root / "meeting"
        self.folder.mkdir()
        self.stack.enter_context(patch.object(output, "OUTPUT_DIR", self.root))
        self.stack.enter_context(patch.object(config, "load_settings", return_value=copy.deepcopy(config.DEFAULTS)))
        self.stack.enter_context(patch.object(app.diarization, "available", return_value=False))
        self.stack.enter_context(patch.object(app.people, "load_people", return_value=[]))
        self.stack.enter_context(patch.object(app.glossary, "apply_corrections", side_effect=lambda t: t))
        self.learn = self.stack.enter_context(patch.object(app.glossary, "learn_from_diff", return_value=0))
        self.job = dict(job_dir=str(self.folder), source="test.wav", duration=30,
                        utterances=[utterance(), utterance("資料を送ります。", 10)])
        pipeline._save_job(self.job)
        self.job = app._load_job("meeting")
        app._job_views(self.job)

    def rows(self, job=None):
        return app._edit_rows(copy.deepcopy(job or self.job))

    def test_edit_history_restores_body_and_document_freshness_and_can_undo_restore(self):
        document_status.save(self.folder, "議事録.md", "以前の議事録", document_status.inputs(self.job))
        rows = self.rows()
        rows[0][3] = "開始は11時です。"
        changed = app.save_utterance_edits(rows, self.job)[1]
        self.assertEqual(self.job["utterances"][0]["text"], "開始は10時です。")
        self.assertEqual(document_status.status(changed, "議事録.md"), "stale")
        version = job_history.choices(changed)[0][1]
        restored = app.restore_edits(changed, version)[1]
        self.assertEqual(restored["utterances"], self.job["utterances"])
        self.assertEqual(document_status.status(restored, "議事録.md"), "current")
        redo = app.restore_edits(restored, job_history.choices(restored)[0][1])[1]
        self.assertEqual(redo["utterances"][0]["text"], "開始は11時です。")

    def test_old_screen_cannot_overwrite_newer_save_or_train_dictionary(self):
        rows = self.rows()
        rows[0][3] = "開始は11時です。"
        app.save_utterance_edits(rows, self.job)
        self.learn.reset_mock()
        rows[0][3] = "開始は12時です。"
        with self.assertRaises(app.gr.Error):
            app.save_utterance_edits(rows, self.job)
        self.learn.assert_not_called()
        self.assertEqual(app._load_job("meeting")["utterances"][0]["text"], "開始は11時です。")
        self.assertTrue(app.workflow_controls(self.job)[4]["interactive"])
        refreshed = app.discard_edits(self.job)[1]
        self.assertEqual(refreshed["utterances"][0]["text"], "開始は11時です。")
        app._require_saved(refreshed)

    def test_manual_correction_clears_pending_ai_review(self):
        self.job["utterances"][0]["proofread_pending"] = True
        pipeline._save_job(self.job)
        self.job = app._load_job("meeting")
        rows = self.rows()
        rows[0][3] = "開始は11時です。"
        saved = app.save_utterance_edits(rows, self.job)[1]
        self.assertNotIn("proofread_pending", saved["utterances"][0])
        self.assertTrue(saved["utterances"][0]["human"])

    def test_invalid_or_missing_rows_do_not_delete_utterances(self):
        valid = self.rows()
        for rows in [valid[:1], [valid[0], valid[0]], [[1.5, "", "", "bad"], valid[1]],
                     [["nan", "", "", "bad"], valid[1]]]:
            with self.subTest(rows=rows), self.assertRaises(app.gr.Error):
                app.save_utterance_edits(rows, self.job)
        self.learn.assert_not_called()
        self.assertEqual(len(app._load_job("meeting")["utterances"]), 2)

    def test_failed_save_preserves_session_and_does_not_train(self):
        rows = self.rows()
        rows[0][3] = "開始は11時です。"
        with patch.object(app, "_write_job", side_effect=OSError("disk full")), self.assertRaises(app.gr.Error):
            app.save_utterance_edits(rows, self.job)
        self.assertEqual(self.job["utterances"][0]["text"], "開始は10時です。")
        self.learn.assert_not_called()

    def test_failed_final_publication_restores_derived_files_and_embeddings(self):
        previous = {path.name: path.read_bytes() for path in self.folder.iterdir() if path.is_file()}
        rows = self.rows()
        rows[0][3] = "開始は11時です。"
        original = output.save_json

        def fail_publish(folder, name, data):
            if name == "utterances.json":
                raise OSError("publication failed")
            return original(folder, name, data)

        with patch.object(output, "save_json", side_effect=fail_publish), self.assertRaises(app.gr.Error):
            app.save_utterance_edits(rows, self.job)
        for name, content in previous.items():
            self.assertEqual((self.folder / name).read_bytes(), content)
        self.assertEqual(len(job_history.choices(self.job)), 1)
        self.learn.assert_not_called()

    def test_unsaved_edit_blocks_navigation_and_ai_then_discard_reloads(self):
        dirty = app.mark_unsaved(self.job, "text")
        with self.assertRaises(app.gr.Error):
            app.open_job("meeting", dirty)
        with patch.object(minutes, "generate_insights") as generate, self.assertRaises(app.gr.Error):
            app.make_insights(dirty, progress=Mock())
        generate.assert_not_called()
        discarded = app.discard_edits(dirty)[1]
        self.assertNotIn("_dirty", discarded)

    def test_empty_generation_does_not_overwrite_previous_document(self):
        document_status.save(self.folder, "要約と考察.md", "以前の内容", document_status.inputs(self.job))
        with patch.object(app, "_llm_ready", return_value=""), \
                patch.object(minutes, "generate_insights", return_value="  "):
            result = app.make_insights(self.job, progress=Mock())
        self.assertIn("作れませんでした", result[0])
        self.assertEqual((self.folder / "要約と考察.md").read_text(encoding="utf-8"), "以前の内容")

    def test_empty_transcript_never_generates_a_document(self):
        empty = copy.deepcopy(self.job)
        empty["utterances"] = []
        with patch.object(minutes, "generate_insights") as generate:
            result = app.make_insights(empty, progress=Mock())
        self.assertIn("発言がない", result[0])
        generate.assert_not_called()

    def test_rediarization_preserves_human_ai_and_legacy_edits_and_timing(self):
        for flag in ("human", "proofread", "legacy"):
            with self.subTest(flag=flag):
                u = utterance("人が確認した本文です。")
                u[flag] = True
                u["words"] = [dict(start=0, end=1, word="元", prob=.9),
                              dict(start=1, end=3, word="の文", prob=.9)]
                result = U.relabel_preserving_edits([u], {0: ["話者2", "話者3"]})
                self.assertEqual(result[0], dict(u, speaker="話者3"))
                self.assertEqual(u["speaker"], "話者1")

    def test_reflow_keeps_utterances_without_word_timestamps(self):
        first = utterance("通常の文。")
        first["words"] = [dict(start=0, end=3, word="通常の文。", prob=1)]
        last = utterance("単語情報のない文。", 10)
        self.assertIn(last, U.reflow([first, last]))

    def test_noise_transfer_selects_audio_and_disables_second_denoise(self):
        wav = self.root / "reduced.wav"
        wav.write_bytes(b"fixture")
        result = app.use_reduced_audio(str(wav), self.job)
        self.assertEqual(result[0], [str(wav)])
        self.assertFalse(result[1])
        self.assertEqual(result[2]["selected"], "transcribe")

    def test_timestamp_player_uses_keyboard_operable_button_and_escapes_text(self):
        html = U.render_html([utterance("<script>bad</script>")], "/audio.m4a")
        self.assertIn('<button class="st-time" type="button"', html)
        self.assertIn("aria-label=", html)
        self.assertNotIn("<script>", html)

    def test_search_avoids_duplicate_transcripts_and_opens_selected_document(self):
        document_status.save(self.folder, "議事録.md", "開始は10時です。", document_status.inputs(self.job))
        with patch.object(search, "OUTPUT_DIR", self.root):
            rows = search.search_outputs("開始")
        self.assertNotIn("transcript.md", [r[1] for r in rows])
        index = next(i for i, row in enumerate(rows) if row[1] == "議事録.md")
        result = app.open_search_result(rows, self.job, SimpleNamespace(index=[index, 3]))
        self.assertEqual(result[1]["job_dir"], str(self.folder))
        self.assertEqual(result[-1]["selected"], "minutes")
        self.assertEqual(result[-2]["selected"], "transcribe")


class ProofreadingTest(unittest.TestCase):
    def test_missing_rejected_duplicate_lines_are_not_marked_complete(self):
        rows = [utterance(), utterance("長い発言を削除せずにそのまま残してください。", 4), utterance("これは最後の行です。", 8)]
        with patch.object(minutes, "_load_prompt", return_value="prompt"), \
                patch.object(minutes, "_context_block", return_value=""), \
                patch.object(minutes.llm, "generate", return_value="[1] 話者1: 開始は11時です。\n[1] 重複の回答\n[2] はい"):
            done, changed = minutes.proofread_utterances({}, rows)
        self.assertEqual((done, changed), (1, 1))
        self.assertTrue(rows[0]["proofread"])
        for row in rows[1:]:
            self.assertFalse(row.get("proofread"))
            self.assertTrue(row["proofread_pending"])

    def test_redo_rechecks_ai_edits_but_preserves_human_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            ai = utterance("開始は11時です。")
            ai.update(proofread=True, words=[dict(start=0, end=3, word="開始は10時です。", prob=1)])
            human = dict(utterance("確認済みです。", 5), human=True)
            job = dict(job_dir=str(folder), source="test.wav", utterances=[ai, human])
            with patch.object(app, "_llm_ready", return_value=""), \
                    patch.object(app.people, "load_people", return_value=[]), \
                    patch.object(app.glossary, "apply_corrections", side_effect=lambda t: t), \
                    patch.object(app, "_save_job"), \
                    patch.object(minutes, "_load_prompt", return_value="prompt"), \
                    patch.object(minutes, "_context_block", return_value=""), \
                    patch.object(minutes.llm, "generate", return_value="[1] 話者1: 開始は12時です。") as generate:
                result = app.proofread_job(job, True, progress=Mock())
            generate.assert_called_once()
            self.assertEqual(result[1]["utterances"][0]["text"], "開始は12時です。")
            self.assertEqual(result[1]["utterances"][1], human)
            self.assertEqual(job["utterances"][0]["text"], "開始は11時です。")


class ExportTablesTest(unittest.TestCase):
    def test_word_contains_real_table_and_pdf_handles_table_across_pages(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            md = folder / "議事録.md"
            md.write_text("# 議事録\n\n| 担当 | 内容 | 期限 |\n| --- | --- | --- |\n"
                          + "".join(f"| 担当A | **資料{i}**を送付 | 来週 |\n" for i in range(90)), encoding="utf-8")
            export.md_to_docx(str(md), str(folder / "out.docx"))
            document = Document(folder / "out.docx")
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.tables[0].rows), 91)
            self.assertEqual(document.tables[0].cell(1, 1).text, "資料0を送付")
            export.md_to_pdf(str(md), str(folder / "out.pdf"))
            self.assertTrue((folder / "out.pdf").read_bytes().startswith(b"%PDF"))

    def test_table_parser_keeps_escaped_pipes_and_extra_cells(self):
        blocks = list(export._blocks("| A | B |\n| --- | :---: |\n| x\\|y | z | extra |"))
        self.assertEqual(blocks, [("table", [["A", "B", ""], ["x|y", "z", "extra"]])])


if __name__ == "__main__":
    unittest.main()
